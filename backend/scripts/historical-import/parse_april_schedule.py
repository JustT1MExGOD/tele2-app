import openpyxl, csv, datetime, sys

SRC = r"C:\Users\shywh\Desktop\tele2-app\шнейне пэпэ апрель.xlsx"
OUT = "source/t2_april_schedule.csv"

ALIASES = {
    "Степанов Алекскей Юрьевич": "Степанов Алексей Юрьевич",
}

# store rule table confirmed by user
def classify(code):
    if code is None:
        return None
    c = str(code).strip()
    if c in ("вых", "отпуск", "отпуску"):
        return ("day_off", None, None, None)
    if c == "9-21":
        return ("working", "kalinina2", "09:00", "21:00")
    if c == "9-17":  # normalizes to 9-21 Kalinina 2 per confirmed rule
        return ("working", "kalinina2", "09:00", "21:00")
    if c == "10-21":
        return ("working", "kosmonavtov", "10:00", "21:00")
    return ("unknown", None, None, None)

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb.worksheets[0]

date_row = 8
dates = []
for col in range(4, 34):
    v = ws.cell(date_row, col).value
    if isinstance(v, datetime.datetime):
        dates.append((col, v.date().isoformat()))

employee_rows = [10, 12, 14, 16, 18, 20]

rows = []
unknown = []
for r in employee_rows:
    name_raw = ws.cell(r, 1).value
    if not name_raw:
        continue
    name = str(name_raw).strip()
    name = ALIASES.get(name, name)
    for col, date_iso in dates:
        code = ws.cell(r, col).value
        result = classify(code)
        if result is None:
            continue
        status, store, start, end = result
        if status == "unknown":
            unknown.append((name, date_iso, code))
            continue
        rows.append([name, date_iso, status, store or "", start or "", end or "", str(code)])

with open(OUT, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["employee_name", "date", "status", "store_id", "shift_start", "shift_end", "raw_code"])
    w.writerows(rows)

print(f"wrote {len(rows)} rows, {len(unknown)} unknown codes", file=sys.stderr)
if unknown:
    print(unknown, file=sys.stderr)
