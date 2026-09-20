# -*- coding: utf-8 -*-
import openpyxl, csv, datetime, sys

SRC_DIR = r"C:\Users\shywh\Desktop\tele2-app"
FILES = [
    ("Планы и график ноябрь 2025.xlsx", "2025-11"),
    ("планы и график Декабрь 2025.xlsx", "2025-12"),
    ("планы и график январь 2026.xlsx", "2026-01"),
    ("планы и график МАРТ 2026.xlsx", "2026-03"),
]
OUT = "source/t2_new_daily_facts.csv"

ALIASES = {
    "Степанов Алекскей Юрьевич": "Степанов Алексей Юрьевич",
}

EXCLUDED_METRICS = {"MNP-факт"}


def find_sheet(wb, needle):
    for name in wb.sheetnames:
        if needle in name.lower():
            return name
    raise KeyError(f"no sheet matching {needle!r} in {wb.sheetnames}")


def build_surname_map(wb):
    """Планы дневные rows use 'SURNAME Firstname' only; Общее план has full
    'Surname Firstname Patronymic'. Resolve via surname (first token), which
    is unique within one file's roster."""
    ws = wb["Общее план"]
    m = {}
    for r in range(2, 9):
        name_raw = ws.cell(r, 1).value
        if not name_raw:
            continue
        name = str(name_raw).split("\n")[0].strip()
        if name == "Общее":
            continue
        surname = name.split()[0].upper()
        m[surname] = ALIASES.get(name, name)
    return m


def process_file(path, fname, month):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_name = find_sheet(wb, "дневны")
    ws = wb[sheet_name]
    surname_map = build_surname_map(wb)

    # locate all header rows (rows containing >=3 datetime cells)
    header_rows = []
    for r in range(1, ws.max_row + 1):
        count = 0
        for c in range(1, ws.max_column + 1):
            if isinstance(ws.cell(r, c).value, datetime.datetime):
                count += 1
        if count >= 3:
            header_rows.append(r)
    if not header_rows:
        raise AssertionError(f"no header rows found in {fname}")

    # canonical date columns from the first header row
    dates = []
    hr0 = header_rows[0]
    for c in range(1, ws.max_column + 1):
        v = ws.cell(hr0, c).value
        if isinstance(v, datetime.datetime):
            dates.append((c, v.date().isoformat()))

    rows = []
    unmapped_labels = set()
    for i, hr in enumerate(header_rows):
        next_hr = header_rows[i + 1] if i + 1 < len(header_rows) else ws.max_row + 1
        name_raw = ws.cell(hr + 1, 1).value
        if not name_raw:
            continue
        short_name = str(name_raw).split("\n")[0].strip()
        surname = short_name.split()[0].upper() if short_name.split() else ""
        name = surname_map.get(surname)
        if not name:
            unmapped_labels.add((fname, short_name, "NO_SURNAME_MATCH", "", ""))
            continue
        for r in range(hr + 1, next_hr):
            metric = ws.cell(r, 2).value
            if metric is None:
                continue
            metric = str(metric).strip()
            if metric in EXCLUDED_METRICS:
                continue
            for c, date_iso in dates:
                val = ws.cell(r, c).value
                if val is None:
                    continue
                if isinstance(val, str) and val.strip() == "":
                    continue
                if not isinstance(val, (int, float)):
                    unmapped_labels.add((fname, name, date_iso, metric, repr(val)))
                    continue
                rows.append([
                    "daily_fact", month, date_iso, name, "продавец", "", "", "", "", "",
                    "", "", metric, val, fname, sheet_name, r, c, "OK", "",
                ])

    return rows, unmapped_labels


all_rows = []
all_bad = []
for fname, month in FILES:
    path = f"{SRC_DIR}\\{fname}"
    rows, bad = process_file(path, fname, month)
    print(f"{fname}: {len(rows)} rows, {len(bad)} bad", file=sys.stderr)
    all_rows.extend(rows)
    all_bad.extend(bad)

HEADER = [
    "record_type", "month", "date", "employee_name", "role", "employment_status",
    "store", "schedule_status", "schedule_code_source", "schedule_code_normalized",
    "shift_start", "shift_end", "metric", "value", "source_file", "source_sheet",
    "source_row", "source_column", "validation_status", "notes",
]

with open(OUT, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(HEADER)
    w.writerows(all_rows)

print(f"TOTAL: wrote {len(all_rows)} rows, {len(all_bad)} bad values", file=sys.stderr)
for b in all_bad:
    print(b, file=sys.stderr)
