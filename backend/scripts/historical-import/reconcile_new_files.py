# -*- coding: utf-8 -*-
import openpyxl, csv, sys
from collections import defaultdict

SRC_DIR = r"C:\Users\shywh\Desktop\tele2-app"
FILES = [
    ("Планы и график ноябрь 2025.xlsx", "2025-11"),
    ("планы и график Декабрь 2025.xlsx", "2025-12"),
    ("планы и график январь 2026.xlsx", "2026-01"),
    ("планы и график МАРТ 2026.xlsx", "2026-03"),
]

METRIC_MAP = {
    "Sim": "sim", "SIM": "sim",
    "MNP-заявки": "mnp",
    "ПА": "pa",
    "Combo": "combo", "Комбо": "combo",
    "Телефон": "phones",
    "Аксы": "accessories",
    "ФО": "focus",
    "Доп услуги": "settings",
    "wink": "wink",
    "ШПД": "shpd",
    "Страховки": "insurance",
    "Кредит": "credit_issued",
}

ALIASES = {
    "Степанов Алекскей Юрьевич": "Степанов Алексей Юрьевич",
}


def norm_name(n):
    n = str(n).strip()
    n = n.replace("\n", " ").strip()
    while "  " in n:
        n = n.replace("  ", " ")
    n = ALIASES.get(n, n)
    return n


# load daily facts sums
daily_sums = defaultdict(float)  # (month, name, metric) -> sum
with open("source/t2_new_daily_facts.csv", encoding="utf-8") as f:
    for row in csv.DictReader(f):
        metric = METRIC_MAP.get(row["metric"])
        if not metric:
            continue
        key = (row["month"], norm_name(row["employee_name"]), metric)
        daily_sums[key] += float(row["value"])

mismatches = []
matches = 0
for fname, month in FILES:
    path = f"{SRC_DIR}\\{fname}"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Общее план"]
    header = [ws.cell(1, c).value for c in range(1, 14)]
    for r in range(2, 9):
        name_raw = ws.cell(r, 1).value
        if not name_raw or str(name_raw).strip() in ("Общее",):
            continue
        name = norm_name(name_raw)
        for c in range(2, 14):
            label = header[c - 1]
            if label not in METRIC_MAP:
                continue
            metric = METRIC_MAP[label]
            plan_val = ws.cell(r, c).value
            if plan_val is None:
                plan_val = 0
            plan_val = float(plan_val)
            daily_val = daily_sums.get((month, name, metric), 0.0)
            if abs(plan_val - daily_val) > 0.01:
                mismatches.append((fname, name, metric, plan_val, daily_val))
            else:
                matches += 1

print(f"matches={matches} mismatches={len(mismatches)}", file=sys.stderr)
for m in mismatches:
    print(m, file=sys.stderr)
